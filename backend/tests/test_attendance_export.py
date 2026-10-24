"""Phase 2 F: выгрузка посещаемости в .xlsx (сводка + лист на кружок)."""
from __future__ import annotations

import secrets
from io import BytesIO

from openpyxl import load_workbook

from .test_attendance import _checkin, _fixed_clock, _qr, set_clock, student_id  # noqa: F401
from .test_exception_conflicts_api import _post, admin_client, make_world  # noqa: F401
from .test_session_club_conflicts_api import _second_club, _session

DEVS = {n: n.encode().hex().ljust(40, "0") for n in ("Анна", "Борис", "Вера")}
IDS = {n: student_id(secrets.token_bytes(32), n, "1ИСИП-126") for n in ("Анна", "Борис", "Вера")}
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _checkins(client, sid, day, names, hm="10:02"):
    set_clock(day=day, hm=hm)
    tok = _qr(client, sid, day=day)["token"]
    for n in names:
        assert _checkin(client, IDS[n], sid, tok, dev=DEVS[n]).status_code == 201


def _world(client):
    w = make_world(client)
    other = _post(client, "/api/admin/leaders", {"name": "Другой", "position": "", "contact": ""})
    room2 = _post(client, "/api/admin/rooms", {"name": "202", "building": "Главный", "capacity": 10})
    c2 = _second_club(client, w, name="Театр «Маска»: [студия]/*", leader_id=other["id"])
    s2 = _session(client, c2, room2, 1, "12:00", "13:00").json()
    _checkins(client, w["session"]["id"], "2027-11-01", ["Анна", "Борис"])
    _checkins(client, w["session"]["id"], "2027-11-08", ["Анна"])
    _checkins(client, s2["id"], "2027-11-01", ["Вера"], hm="12:01")
    return w, c2


def test_export_all_clubs(admin_client):
    w, c2 = _world(admin_client)
    r = admin_client.get("/api/admin/attendance/export.xlsx?date_from=2027-10-01&date_to=2027-11-30")
    assert r.status_code == 200 and r.headers["content-type"] == XLSX
    assert "attendance_2027-10-01_2027-11-30.xlsx" in r.headers["content-disposition"]
    wb = load_workbook(BytesIO(r.content))
    assert wb.sheetnames[0] == "Сводка" and len(wb.sheetnames) == 3
    summary = {row[0]: row[1:3] for row in wb["Сводка"].iter_rows(min_row=4, values_only=True)}
    assert summary["Python"] == (2, 3) and summary["Итого"] == (3, 4)
    py = list(wb["Python"].iter_rows(values_only=True))
    assert py[0] == ("Дата", "Время", "Преподаватель", "ID студента", "Отметился в")
    got = [(r[0].date().isoformat(), r[3]) for r in py[1:]]
    assert sorted(got) == sorted([("2027-11-01", IDS["Анна"]), ("2027-11-01", IDS["Борис"]), ("2027-11-08", IDS["Анна"])])
    assert py[1][2]  # имя преподавателя есть
    assert not any(n in str(r) for r in py for n in IDS)  # имён студентов нет
    theatre = [s for s in wb.sheetnames if s.startswith("Театр")][0]
    assert not set("[]:*?/\\") & set(theatre) and len(theatre) <= 31


def test_export_one_club_and_validation(admin_client):
    w, c2 = _world(admin_client)
    r = admin_client.get(f"/api/admin/attendance/export.xlsx?date_from=2027-10-01&date_to=2027-11-30&club_id={c2['id']}")
    wb = load_workbook(BytesIO(r.content))
    assert len(wb.sheetnames) == 2 and wb["Сводка"]["A4"].value.startswith("Театр")
    assert admin_client.get("/api/admin/attendance/export.xlsx?date_from=2027-12-01&date_to=2027-11-01").status_code == 422
    empty = load_workbook(BytesIO(admin_client.get("/api/admin/attendance/export.xlsx?date_from=2020-01-01&date_to=2020-01-31").content))
    assert empty.sheetnames == ["Сводка"]


def test_export_admin_only(client):
    assert client.get("/api/admin/attendance/export.xlsx").status_code == 401
